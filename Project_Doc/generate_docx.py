import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_table(doc, rows, cols, col_widths, headers, data):
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B") # Dark slate
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = 'Calibri'

    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = 'Calibri'
                run.font.color.rgb = RGBColor(51, 65, 85)

    # Set column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

    # Add spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(8)
    return table

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        run.font.name = 'Calibri'
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
        run.font.name = 'Calibri'
    return h

def add_heading_3(doc, text):
    h = doc.add_heading(text, level=3)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    for run in h.runs:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
        run.font.name = 'Calibri'
    return h

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.bold = True
        r_bold.font.size = Pt(10)
        r_bold.font.name = 'Calibri'
        r_bold.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = 'Calibri'
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, bold_title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_title:
        r_b = p.add_run(bold_title + ": ")
        r_b.font.bold = True
        r_b.font.size = Pt(10)
        r_b.font.name = 'Calibri'
        r_b.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = 'Calibri'
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r_t = p.add_run(f"[{title}] ")
    r_t.font.bold = True
    r_t.font.size = Pt(9.5)
    r_t.font.color.rgb = RGBColor(30, 41, 59)
    r_c = p.add_run(text)
    r_c.font.size = Pt(9.5)
    r_c.font.color.rgb = RGBColor(71, 85, 105)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document(output_path):
    doc = docx.Document()

    # Configure Margins (0.75 in on all sides)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("Orangyy Carpels (Studio OS)")
    t_run.font.size = Pt(24)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(234, 88, 12) # Orange 600
    t_run.font.name = 'Calibri'

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    s_run = sub_p.add_run("Complete System Architecture & Implementation Specification — Phase G V1")
    s_run.font.size = Pt(13)
    s_run.font.bold = True
    s_run.font.color.rgb = RGBColor(71, 85, 105)
    s_run.font.name = 'Calibri'

    # Meta table
    meta_headers = ["Document Version", "Release Phase", "Target Database", "Architecture"]
    meta_data = [["1.0 Production", "Phase G V1 (Complete)", "MySQL 8.0 (Unified Local/Server)", "Modular Monolith (Service-Layer)"]]
    create_table(doc, 2, 4, [1.5, 1.6, 1.8, 1.6], meta_headers, meta_data)

    # --- Section 1: Executive Summary & System Overview ---
    add_heading_1(doc, "1. Executive Summary & System Overview")
    add_body_p(doc, "Orangyy Carpels is a centralized enterprise platform engineered specifically for design studios, consulting agencies, and project-based organizations. It integrates team management, client records, project delivery, monthly/daily timesheets, leave management, multi-currency billing, and business reporting into a single high-performance system.")

    add_heading_2(doc, "1.1 Technical Stack Specifications")
    tech_headers = ["Layer", "Technology", "Description & Rules"]
    tech_data = [
        ["Frontend UI", "React 18 + Vite", "Modular Single Page Application with React Router v6"],
        ["Design System", "Tailwind CSS", "Dylan Field (Figma) inspired minimalist, light aesthetic"],
        ["Language", "TypeScript", "Strict type-checking across frontend and backend"],
        ["Backend Server", "Node.js + Fastify", "High-performance modular API server with JSON Schema validation"],
        ["ORM", "Prisma ORM", "Type-safe database abstraction with MySQL connector"],
        ["Database", "MySQL 8.0", "100% unified standard across local and production environments"],
        ["Authentication", "Cookie Sessions", "Signed HTTP-only cookie sessions with salted scrypt hashing"],
        ["External APIs", "open.er-api.com", "Live multi-currency foreign exchange rates against INR"],
        ["Process Manager", "PM2", "Zero-downtime cluster mode and daemon process manager"],
        ["Web Server", "Nginx", "Reverse proxy, SSL termination, and static asset delivery"]
    ]
    create_table(doc, len(tech_data) + 1, 3, [1.5, 1.8, 3.2], tech_headers, tech_data)

    # --- Section 2: Complete Database Architecture ---
    add_heading_1(doc, "2. Complete Database Architecture & Schema Specification")
    add_body_p(doc, "The database is built on MySQL 8.0 using Prisma ORM. It balances relational data integrity with high-performance JSON sub-documents for rate history, monthly budgets, education, and career experience.")

    add_heading_2(doc, "2.1 Model: Employee (HR & Identity)")
    emp_headers = ["Field Name", "Type", "Constraints / Default", "Description"]
    emp_data = [
        ["employeeId", "String", "@id (PK)", "Unique Employee Code (e.g. EMP001)"],
        ["fullName", "String", "Required", "Legal full name"],
        ["designation", "String", 'default("Team Member")', "Professional role title"],
        ["department", "String", 'default("General")', "Department (Design, Engineering, etc.)"],
        ["email", "String", "@unique", "Corporate email address"],
        ["password", "String?", 'default("")', "Salted scrypt password hash"],
        ["phone", "String", "Required", "Primary mobile number"],
        ["costRate", "String", 'default("₹0/hr")', "Internal cost/billing rate per hour"],
        ["capacity", "String", 'default("40 hrs/week")', "Weekly capacity hours standard"],
        ["role", "String", 'default("Employee")', "Super Admin, Project Manager, Employee"],
        ["status", "String", 'default("Active")', "Active or Inactive"],
        ["education", "Json?", "Optional", "Array of { degree, school, year }"],
        ["experience", "Json?", "Optional", "Array of { company, role, period }"],
        ["leaveYear", "Int", "default(2026)", "Active calendar/fiscal leave year"],
        ["casualQuota / Used", "Float", "default(12) / (0)", "Annual Casual Leave allotment / consumed"],
        ["sickQuota / Used", "Float", "default(12) / (0)", "Annual Sick Leave allotment / consumed"],
        ["earnedQuota / Used", "Float", "default(15) / (0)", "Annual Earned Leave allotment / consumed"],
        ["compOffBalance", "Float", "default(0)", "Earned Compensatory Off balance (days)"],
        ["optionalHolidaysQuota / Used", "Int", "default(2) / (0)", "Annual optional holidays allowed / used"],
        ["wfhMonthlyLimit / Used", "Int", "default(2) / (0)", "Monthly Work From Home quota / used"]
    ]
    create_table(doc, len(emp_data) + 1, 4, [1.5, 0.8, 1.6, 2.6], emp_headers, emp_data)

    add_heading_2(doc, "2.2 Model: Client (CRM & Statutory)")
    cli_headers = ["Field Name", "Type", "Constraints / Default", "Description"]
    cli_data = [
        ["id", "String", "@id (PK)", "Client Unique Identifier (CLI-001)"],
        ["name", "String", "Required", "Primary commercial brand name"],
        ["legalName", "String?", 'default("")', "Registered statutory legal name"],
        ["billingCurrency", "String", 'default("USD ($)")', "Base currency for billing"],
        ["defaultBillingType", "String", 'default("T&M")', "T&M, Fixed RC, or Fixed PC"],
        ["dueTime", "String", 'default("30 days")', "Payment credit term (15, 30, 45, 60, 90 days)"],
        ["cin / gst / pan / msme", "String?", 'default("")', "Corporate identification & tax numbers"],
        ["status", "String", 'default("Active")', "Active or Inactive"],
        ["projects", "Project[]", "1:N Cascade", "Associated client projects"]
    ]
    create_table(doc, len(cli_data) + 1, 4, [1.5, 0.8, 1.6, 2.6], cli_headers, cli_data)

    add_heading_2(doc, "2.3 Model: Project (Delivery & Budgets)")
    prj_headers = ["Field Name", "Type", "Constraints / Default", "Description"]
    prj_data = [
        ["id", "String", "@id (PK)", "Project Identifier (PRJ-001)"],
        ["name", "String", "Required", "Project title"],
        ["clientId", "String", "FK -> Client.id", "Parent client account"],
        ["billingType", "String", "Required", "T&M, Fixed RC, or Fixed PC"],
        ["rate / currency", "String", "Required", "Base rate string and currency"],
        ["budgetHours", "Int", "default(0)", "Active monthly budget hours"],
        ["loggedHours", "Int", "default(0)", "Total historical logged hours (auto-synced)"],
        ["managerId / Name", "String?", 'default("")', "Designated Project Manager"],
        ["assignedEmployees", "String?", 'default("")', "Comma-separated list of assigned resource IDs"],
        ["monthlyBudgets", "Json?", "Optional", "Array of { monthYear, budgetHours, isLocked }"],
        ["rateVersions", "Json?", "Optional", "Array of { id, billingType, rateAmount, effectiveStartDate, effectiveEndDate }"],
        ["dailyEntries", "DailyEntry[]", "1:N Cascade", "Daily timesheet records"]
    ]
    create_table(doc, len(prj_data) + 1, 4, [1.5, 0.8, 1.6, 2.6], prj_headers, prj_data)

    add_heading_2(doc, "2.4 Model: DailyTimesheetEntry (Work Logs)")
    dt_headers = ["Field Name", "Type", "Constraints / Default", "Description"]
    dt_data = [
        ["id", "Int", "@id @default(autoincrement())", "Auto-increment primary key"],
        ["projectId", "String", "FK -> Project.id", "Associated project"],
        ["employeeId", "String?", 'default("")', "Employee ID who logged the entry"],
        ["date", "String", "Required (YYYY-MM-DD)", "Date of work performed"],
        ["hours", "Float", "default(0)", "Decimal hours logged (e.g. 4.5)"],
        ["task", "String", 'default("")', "Ideation, Design, AI Design, Validation, Analysis, Approval"],
        ["description", "String", 'default("")', "Detailed task notes"],
        ["isBillable", "Boolean", "default(true)", "Billable vs Non-Billable indicator"],
        ["status", "String", 'default("Draft")', "Draft, Submitted, PM_Approved, Approved"],
        ["@@unique", "Constraint", "[projectId, date, employeeId]", "Prevents duplicate overlapping entries per day"]
    ]
    create_table(doc, len(dt_data) + 1, 4, [1.5, 0.8, 1.6, 2.6], dt_headers, dt_data)

    add_heading_2(doc, "2.5 Model: LeaveRequest (Absences & Approvals)")
    lr_headers = ["Field Name", "Type", "Constraints / Default", "Description"]
    lr_data = [
        ["id", "Int", "@id @default(autoincrement())", "Auto-increment primary key"],
        ["employeeId / Name", "String", "Required", "Employee who submitted the request"],
        ["leaveType", "String", "Required", "Casual Leave, Sick Leave, Earned Leave, Comp-off, WFH, Optional Holiday"],
        ["startDate / endDate", "String", "Required (YYYY-MM-DD)", "Leave duration dates"],
        ["isHalfDay / session", "Boolean / String?", "default(false)", "Half-day flag and session (First Half / Second Half)"],
        ["daysCount", "Float", "Required", "Business days count (0.5, 1.0, 2.0...)"],
        ["pmApproval", "String?", 'default("Pending")', "Pending, Approved, Rejected"],
        ["saApproval", "String?", 'default("Pending")', "Pending, Approved, Rejected"],
        ["status", "String", 'default("Pending")', "Pending_PM, Pending_SA, Approved, Rejected, Cancelled"]
    ]
    create_table(doc, len(lr_data) + 1, 4, [1.5, 0.8, 1.6, 2.6], lr_headers, lr_data)

    add_heading_2(doc, "2.6 Model: ExchangeRate (Forex Engine)")
    fx_headers = ["Field Name", "Type", "Constraints / Default", "Description"]
    fx_data = [
        ["id", "Int", "@id @default(autoincrement())", "Auto-increment primary key"],
        ["currency", "String", "Required (USD, EUR, GBP...)", "3-letter currency code"],
        ["rateToINR", "Float", "Required", "Conversion multiplier to Indian Rupee (INR)"],
        ["monthYear", "String", "Required (YYYY-MM)", "Active monthly partition key"],
        ["isLocked", "Boolean", "default(false)", "Locked state once month closes"],
        ["source", "String", 'default("online")', "open.er-api.com or studio-default"]
    ]
    create_table(doc, len(fx_data) + 1, 4, [1.5, 0.8, 1.6, 2.6], fx_headers, fx_data)

    # --- Section 3: Core Modules & Business Logic ---
    doc.add_page_break()
    add_heading_1(doc, "3. Core Modules & Business Logic")

    add_heading_2(doc, "3.1 Authentication & Role Security Module")
    add_bullet(doc, "Session Management", "Stateful signed HTTP-only cookie (sessionId) verified via Fastify preHandler hook.")
    add_bullet(doc, "Password Hashing", "Salted scrypt hashing using crypto.scryptSync with backward compatibility for plain credentials.")
    add_bullet(doc, "Role Hierarchy", "Super Admin (full governance), Project Manager (team approvals & delivery), Employee (timesheet logging & leave applications).")

    add_heading_2(doc, "3.2 Timesheet & Time Tracking Engine")
    add_bullet(doc, "Project-Centric Monthly Grid", "Renders full month days (01 to 31) with weekend indicators and task selectors.")
    add_bullet(doc, "Debounced Auto-Save", "Inline updates are automatically persisted via POST /api/timesheets/daily with a 500ms debounce.")
    add_bullet(doc, "Dual-Stage Approval Workflow", "Draft -> Submitted -> PM_Approved -> Approved. PM approval is blocked if any assigned resource is pending (Partially_Submitted prevention).")
    add_bullet(doc, "Reopen / Rework", "PMs and Super Admins can unlock submitted sheets with feedback, reverting status to Draft.")
    add_bullet(doc, "Auto Hours Synchronization", "TimesheetService.syncProjectLoggedHours aggregates all entries and syncs Project.loggedHours on every change.")

    add_heading_2(doc, "3.3 Leave & Attendance Engine")
    add_bullet(doc, "Accruals & Quotas", "Casual Leave (12d), Sick Leave (12d), Earned Leave (15d, 1.25/mo), WFH (2d/mo limit), Optional Holidays (2d/yr).")
    add_bullet(doc, "Dual Approval Chain", "Employee applications require PM Approval first, then escalate to Super Admin Approval. PM applications go directly to Super Admin.")
    add_bullet(doc, "Atomic Balance Deductions", "Upon final approval, employee balance counters (casualUsed, sickUsed, compOffBalance) update automatically.")
    add_bullet(doc, "Comp-Off Mechanism", "Employees log extra hours worked on weekends/holidays, which credits their Comp-Off balance once approved.")

    add_heading_2(doc, "3.4 Billing & Multi-Currency Engine")
    add_bullet(doc, "T&M Billing Formula", "T&M Revenue (INR) = Sum(Logged Hours * Hourly Rate * ExchangeRateToINR)")
    add_bullet(doc, "Fixed RC Billing Formula", "Fixed RC Revenue (INR) = Sum(Monthly Resource Rate * ExchangeRateToINR)")
    add_bullet(doc, "Fixed PC Billing Formula", "Fixed PC Revenue (INR) = Sum(Project Cost * ExchangeRateToINR)")
    add_bullet(doc, "Live Forex Sync", "Live rates fetched from open.er-api.com and cached by month-year partition with fallback defaults.")
    add_bullet(doc, "Fiscal Year Partitioning", "Calculates revenue across Indian Fiscal Year (April 1 to March 31).")

    # --- Section 4: RBAC Permissions Matrix ---
    add_heading_1(doc, "4. Role-Based Access Control (RBAC) Matrix")
    rbac_headers = ["Feature / Action", "Employee", "Project Manager", "Super Admin"]
    rbac_data = [
        ["View Dashboard (Personal KPIs)", "YES", "YES", "YES"],
        ["View Studio Financial Revenue (INR)", "NO", "YES", "YES"],
        ["Log & Submit Daily Timesheets", "YES", "YES", "YES"],
        ["Approve Subordinate Timesheets (PM Stage)", "NO", "YES", "YES"],
        ["Final Lock Timesheets (SA Stage)", "NO", "NO", "YES"],
        ["Reopen Submitted Timesheet for Rework", "NO", "YES", "YES"],
        ["View Employee & Client Registries", "NO", "YES", "YES"],
        ["Create / Edit / Delete Employees", "NO", "NO", "YES"],
        ["Create / Edit Clients & Projects", "NO", "NO", "YES"],
        ["Edit Project Rate Versions & Rates", "NO", "NO", "YES"],
        ["Set Project Monthly Budget Hours", "NO", "YES", "YES"],
        ["Apply for Leave & Comp-Off", "YES", "YES", "YES"],
        ["Approve Leave Requests", "NO", "YES (PM Stage)", "YES (Final)"],
        ["Configure Studio Settings & Policies", "NO", "NO", "YES"]
    ]
    create_table(doc, len(rbac_data) + 1, 4, [2.5, 1.2, 1.4, 1.4], rbac_headers, rbac_data)

    # --- Section 5: API Catalog ---
    doc.add_page_break()
    add_heading_1(doc, "5. REST API Route Catalog & Endpoints")
    api_headers = ["Module", "Method", "Endpoint", "Auth / Role", "Description"]
    api_data = [
        ["Auth", "POST", "/api/auth/login", "Public", "Authenticates credentials & issues session cookie"],
        ["Auth", "POST", "/api/auth/logout", "Session", "Destroys active session cookie"],
        ["Auth", "GET", "/api/auth/me", "Session", "Returns current authenticated user session"],
        ["Timesheet", "GET", "/api/timesheets/summary", "All", "List projects with budget vs logged hours"],
        ["Timesheet", "GET", "/api/timesheets/daily", "All", "Get 31-day entries & lock status for a project"],
        ["Timesheet", "POST", "/api/timesheets/daily", "Assigned", "Debounced upsert of daily timesheet entries"],
        ["Timesheet", "POST", "/api/timesheets/approve", "PM / SA", "Advances status to PM_Approved or Approved"],
        ["Timesheet", "POST", "/api/timesheets/reopen", "PM / SA", "Reverts timesheet to Draft for rework"],
        ["Registry", "GET/POST", "/api/registries/employees", "PM / SA (W: SA)", "Employee registry CRUD operations"],
        ["Registry", "GET/POST", "/api/registries/clients", "PM / SA (W: SA)", "Client CRM registry CRUD operations"],
        ["Registry", "GET/POST", "/api/registries/projects", "PM / SA (W: SA)", "Project registry CRUD operations"],
        ["Billing", "GET", "/api/billing/summary", "PM / SA", "Aggregated revenue in INR and project breakdown"],
        ["Billing", "GET/POST", "/api/billing/rates", "PM / SA (W: SA)", "Forex exchange rates and live sync"],
        ["Billing", "GET/POST", "/api/billing/projects/:id/rate-versions", "SA", "Project effective rate version history"],
        ["Billing", "GET/POST", "/api/billing/projects/:id/monthly-budgets", "PM / SA", "Monthly budget hours management"],
        ["Leaves", "GET", "/api/leaves/balance", "All", "Get caller leave balances and quotas"],
        ["Leaves", "POST", "/api/leaves/apply", "All", "Apply for Leave, WFH, or Comp-off"],
        ["Leaves", "POST", "/api/leaves/requests/:id/action", "PM / SA", "Approve or Reject leave request"],
        ["Leaves", "GET", "/api/leaves/attendance-matrix", "All", "Team availability calendar and holidays"]
    ]
    create_table(doc, len(api_data) + 1, 5, [1.0, 0.8, 2.1, 1.1, 1.5], api_headers, api_data)

    # --- Section 6: UI/UX & Deployment ---
    add_heading_1(doc, "6. Frontend Architecture & Deployment Operations")
    add_heading_2(doc, "6.1 UI Design System (Dylan Field / Figma Inspired)")
    add_bullet(doc, "Visual Tone", "Subdued borders (#E5E7EB), light studio backgrounds (#F9FAFB), high-contrast typography, and functional layouts.")
    add_bullet(doc, "Zero Heavy Frameworks", "Pure Tailwind CSS with lightweight custom primitives in frontend/src/components/ui/.")
    add_bullet(doc, "Live Auto-Save Feedback", "Visual state pills (Saving... / All changes saved) provide instant user clarity.")

    add_heading_2(doc, "6.2 Deployment & Environment Maintenance")
    add_bullet(doc, "Process Management (PM2)", "Cluster mode execution with zero downtime via ecosystem.config.js on port 5001.")
    add_bullet(doc, "Web Server (Nginx)", "Reverse proxy routes /api/* to Fastify (5001) and serves pre-built Vite assets with SPA fallback.")
    add_bullet(doc, "Unified Database Sync", "Automated shell scripts (sync-db-to-fortest.sh, sync-db-from-fortest.sh, deploy-now.sh) ensure 100% database schema parity.")

    add_callout(doc, "This specification document represents the exact, frozen production build of Orangyy Carpels Phase G V1.", title="PRODUCTION VERIFICATION")

    # Save Document
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_document("/Users/ckarunkumar/OrangyyDesign/OrangyyCarpels/Project_Doc/Orangyy_Carpels_System_Documentation.docx")
